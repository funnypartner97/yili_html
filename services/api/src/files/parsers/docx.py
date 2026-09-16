from pathlib import Path

from docx import Document

from src.files.parsers.base import MAX_CELLS, MAX_ROWS, MAX_TEXT, finalize, inspect_file, invalid, limit
from src.files.types import ParsedSource, Segment, SourceLocator, SourceTable


class DocxParser:
    def parse(self, path: Path) -> ParsedSource:
        if inspect_file(path, path.name) != 'docx':
            raise invalid()
        document = Document(path)
        if len(document.paragraphs) > 10_000 or len(document.tables) > 100:
            raise limit()
        segments = []
        total = 0
        for index, paragraph in enumerate(document.paragraphs, 1):
            total += len(paragraph.text)
            if total > MAX_TEXT:
                raise limit()
            segments.append(Segment(text=paragraph.text,
                role='heading' if paragraph.style and paragraph.style.name.startswith('Heading') else 'text',
                locator=SourceLocator(paragraph=index)))
        tables = []
        for index, table in enumerate(document.tables, 1):
            if len(table.rows) > MAX_ROWS or len(table.rows) * len(table.columns) > MAX_CELLS:
                raise limit()
            # Paragraph index of the table's first preceding body paragraph + 1.
            anchor = 1 + sum(child.tag.endswith('}p') for child in table._element.itersiblings(preceding=True))
            tables.append(SourceTable(name=f'Table {index}', locator=SourceLocator(paragraph=anchor),
                                      rows=[[cell.text for cell in row.cells] for row in table.rows]))
        return finalize(ParsedSource(kind='docx', title=path.stem, segments=segments, tables=tables))
