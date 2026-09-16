from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from src.files.parsers.base import (MAX_CELL_TEXT, MAX_CELLS, MAX_COLUMNS, MAX_ROWS, MAX_TEXT,
                                    finalize, inspect_file, invalid, limit, read_csv)
from src.files.types import FormulaCell, ParsedSource, SourceLocator, SourceTable


def cell_value(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if not isinstance(value, (str, int, float, bool, type(None))):
        raise invalid()
    if isinstance(value, str) and len(value) > MAX_CELL_TEXT:
        raise limit()
    return value


def formula_cell(value, cache, locator):
    # openpyxl represents these OOXML records as objects; never stringify the
    # Python object or pass it to scalar coercion, and never evaluate anything.
    if isinstance(value, ArrayFormula):
        return FormulaCell(formula=cell_value(value.text), formula_kind='array',
                           formula_range=value.ref, cached_value=cache, locator=locator)
    if isinstance(value, DataTableFormula):
        return FormulaCell(formula_kind='dataTable', formula_range=value.ref,
                           attributes=dict(value), cached_value=cache, locator=locator)
    if isinstance(value, str):
        return FormulaCell(formula=cell_value(value), cached_value=cache, locator=locator)
    raise invalid()


class SpreadsheetParser:
    def parse(self, path: Path) -> ParsedSource:
        kind = inspect_file(path, path.name)
        if kind == 'csv':
            return finalize(ParsedSource(kind='csv', title=path.stem, tables=[SourceTable(
                name=path.stem, locator=SourceLocator(sheet=path.stem), rows=read_csv(path))]))
        if kind != 'xlsx':
            raise invalid()
        formulas_book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            values_book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                if len(formulas_book.worksheets) > 100:
                    raise limit()
                tables = []
                total_cells = total_text = 0
                for sheet in formulas_book:
                    if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                        raise limit()
                    # Ignore forged dimensions but enforce actual streamed row/cell counts.
                    sheet.reset_dimensions()
                    values_sheet = values_book[sheet.title]
                    values_sheet.reset_dimensions()
                    rows, formulas = [], []
                    cached_rows = values_sheet.iter_rows()
                    for row_number, row in enumerate(sheet.iter_rows(), 1):
                        cached = next(cached_rows, ())
                        total_cells += len(row)
                        if row_number > MAX_ROWS or len(row) > MAX_COLUMNS or total_cells > MAX_CELLS:
                            raise limit()
                        values = []
                        for column, cell in enumerate(row, 1):
                            if cell.data_type == 'f':
                                cache = cell_value(cached[column - 1].value) if column <= len(cached) else None
                                formula = formula_cell(cell.value, cache,
                                    SourceLocator(sheet=sheet.title, row=row_number, column=column))
                                formulas.append(formula)
                                value = formula.formula if formula.formula is not None else cache
                            else:
                                value = cell_value(cell.value)
                            total_text += len(str(value)) if value is not None else 0
                            if total_text > MAX_TEXT:
                                raise limit()
                            values.append(value)
                        rows.append(values)
                    tables.append(SourceTable(name=sheet.title, locator=SourceLocator(sheet=sheet.title), rows=rows, formulas=formulas))
                return finalize(ParsedSource(kind='xlsx', title=path.stem, tables=tables))
            finally:
                values_book.close()
        finally:
            formulas_book.close()
