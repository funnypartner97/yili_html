from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.files.parsers.base import MAX_CELLS, MAX_TEXT, finalize, inspect_file, invalid, limit
from src.files.types import ParsedSource, Segment, SourceLocator, SourceTable


def text_shapes(shapes, depth=0):
    if depth > 20:
        raise limit()
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from text_shapes(shape.shapes, depth + 1)
        else:
            yield shape


class PptxParser:
    def parse(self, path: Path) -> ParsedSource:
        if inspect_file(path, path.name) != 'pptx':
            raise invalid()
        document = Presentation(path)
        if len(document.slides) > 1_000:
            raise limit()
        segments, tables = [], []
        total = shapes_count = 0
        for index, slide in enumerate(document.slides, 1):
            locator = SourceLocator(slide=index)
            for shape in text_shapes(slide.shapes):
                shapes_count += 1
                if shapes_count > 10_000:
                    raise limit()
                if shape.has_text_frame:
                    total += len(shape.text)
                    if total > MAX_TEXT:
                        raise limit()
                    if shape.text:
                        segments.append(Segment(text=shape.text, locator=locator))
                if shape.has_table:
                    table = shape.table
                    if len(tables) >= 100 or len(table.rows) * len(table.columns) > MAX_CELLS:
                        raise limit()
                    tables.append(SourceTable(name=f'Slide {index} table {len(tables) + 1}', locator=locator,
                        rows=[[cell.text for cell in row.cells] for row in table.rows]))
        return finalize(ParsedSource(kind='pptx', title=path.stem, segments=segments, tables=tables,
                                     metadata={'slideCount': len(document.slides)}))
