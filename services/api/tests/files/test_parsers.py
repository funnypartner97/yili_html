from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pydantic import ValidationError

from src.core.errors import DomainError
from src.files.parsers.pdf import PdfParser
from src.files.parsers.docx import DocxParser
from src.files.parsers.pptx import PptxParser
from src.files.parsers.spreadsheet import SpreadsheetParser
from src.files.parsers.image import ImageParser
from src.files.parsers.base import inspect_file
from src.files.types import ParsedSource


def test_xlsx_extracts_sheet_cells_and_formula_with_cached_value(tmp_path):
    path = tmp_path / 'sales.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Sales'
    sheet.append(['Month', 'Revenue'])
    sheet.append(['Jan', 120])
    sheet.append(['Total', '=SUM(B2:B2)'])
    workbook.save(path)
    # Excel cache is independent of the formula and must not be calculated here.
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members['xl/worksheets/sheet1.xml'] = members['xl/worksheets/sheet1.xml'].replace(
        b'<f>SUM(B2:B2)</f><v></v>', b'<f>SUM(B2:B2)</f><v>120</v>')
    with ZipFile(path, 'w') as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    parsed = SpreadsheetParser().parse(path)
    assert parsed.kind == 'xlsx'
    assert parsed.tables[0].name == 'Sales'
    assert parsed.tables[0].rows[1] == ['Jan', 120]
    cell = parsed.tables[0].formulas[0]
    assert cell.formula == '=SUM(B2:B2)'
    assert cell.cached_value == 120
    assert cell.locator.model_dump(exclude_none=True) == {'sheet': 'Sales', 'row': 3, 'column': 2}
    assert parsed.tables[0].locator.sheet == 'Sales'


def test_pdf_extracts_page_text(tmp_path):
    path = tmp_path / 'report.pdf'
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), 'Quarterly report')
        document.save(path)
    parsed = PdfParser().parse(path)
    assert parsed.kind == 'pdf'
    assert parsed.segments[0].text.strip() == 'Quarterly report'
    assert parsed.segments[0].locator.page == 1


def test_docx_extracts_headings_with_paragraph_locators(tmp_path):
    path = tmp_path / 'report.docx'
    document = Document()
    document.add_heading('Revenue', 1)
    document.add_paragraph('Revenue grew.')
    document.save(path)
    parsed = DocxParser().parse(path)
    assert parsed.kind == 'docx'
    assert parsed.segments[0].role == 'heading'
    assert parsed.segments[0].text == 'Revenue'
    assert parsed.segments[1].locator.paragraph == 2


def test_pptx_extracts_slide_text(tmp_path):
    path = tmp_path / 'report.pptx'
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[0])
    slide.shapes.title.text = 'Revenue'
    document.save(path)
    parsed = PptxParser().parse(path)
    assert parsed.kind == 'pptx'
    assert parsed.segments[0].text == 'Revenue'
    assert parsed.segments[0].locator.slide == 1


def test_csv_preserves_bom_cells_and_formula_as_inert_text(tmp_path):
    path = tmp_path / 'sales.csv'
    path.write_bytes(b'\xef\xbb\xbfMonth;Revenue\r\nJan;120\r\nTotal;=SUM(B2:B2)\r\n')
    parsed = SpreadsheetParser().parse(path)
    assert parsed.kind == 'csv'
    assert parsed.tables[0].rows == [['Month', 'Revenue'], ['Jan', '120'], ['Total', '=SUM(B2:B2)']]
    assert parsed.tables[0].locator.sheet == 'sales'


def test_csv_keeps_blank_records_so_locators_do_not_shift(tmp_path):
    path = tmp_path / 'sales.csv'
    path.write_bytes(b'Month,Revenue\n\nJan,120\n')
    table = SpreadsheetParser().parse(path).tables[0]
    assert table.rows == [['Month', 'Revenue'], [], ['Jan', '120']]
    assert table.cell_locator(2, 1).row == 3


@pytest.mark.parametrize(('extension', 'kind'), [('png', 'png'), ('jpg', 'jpeg')])
def test_image_extracts_dimensions_without_ocr(tmp_path, extension, kind):
    path = tmp_path / f'image.{extension}'
    Image.new('RGB', (24, 12)).save(path)
    parsed = ImageParser().parse(path)
    assert parsed.kind == kind
    assert parsed.images[0].width == 24
    assert parsed.images[0].height == 12
    assert parsed.segments == []


def test_parsed_source_rejects_unknown_fields_and_coercion():
    with pytest.raises(ValidationError):
        ParsedSource(kind='csv', title='test', unexpected=True)
    with pytest.raises(ValidationError):
        ParsedSource(kind='csv', title=42)


@pytest.mark.parametrize('payload', [b'not-a-pdf', b'<html>hello</html>', b'\xff\x00binary'])
def test_invalid_signature_is_rejected(tmp_path, payload):
    path = tmp_path / 'test.pdf'
    path.write_bytes(payload)
    with pytest.raises(DomainError, match='supported'):
        inspect_file(path, 'test.pdf')


@pytest.mark.parametrize('name', ['../evil.xml', '/evil.xml', 'xl/vbaProject.bin', 'xl/activeX/control.xml'])
def test_unsafe_zip_members_are_rejected(tmp_path, name):
    path = tmp_path / 'test.xlsx'
    workbook = Workbook()
    workbook.save(path)
    with ZipFile(path, 'a') as archive:
        archive.writestr(name, b'evil')
    with pytest.raises(DomainError):
        SpreadsheetParser().parse(path)


def test_compression_bomb_is_rejected_before_parsing(tmp_path):
    path = tmp_path / 'test.xlsx'
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', b'0' * 2_000_000)
    with pytest.raises(DomainError):
        inspect_file(path, 'test.xlsx')


def test_xml_entities_are_rejected(tmp_path):
    path = tmp_path / 'test.docx'
    document = Document()
    document.save(path)
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members['word/document.xml'] = b'<!DOCTYPE x [<!ENTITY x SYSTEM "http://127.0.0.1/secret">]><x>&x;</x>'
    with ZipFile(path, 'w') as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    with pytest.raises(DomainError):
        DocxParser().parse(path)


def test_csv_rejects_ragged_rows_and_overlarge_cells(tmp_path):
    path = tmp_path / 'sales.csv'
    path.write_text('a,b\n1,2,3\n', encoding='utf-8')
    with pytest.raises(DomainError):
        SpreadsheetParser().parse(path)
    path.write_text('a,b\n' + 'x' * 40_000 + ',2\n', encoding='utf-8')
    with pytest.raises(DomainError):
        SpreadsheetParser().parse(path)


@pytest.mark.parametrize('coordinate', ['A1048576', 'XFD1'])
def test_sparse_xlsx_is_rejected_by_preflight_before_row_allocation(tmp_path, coordinate):
    path = tmp_path / 'sparse.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = 'value'
    workbook.save(path)
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members['xl/worksheets/sheet1.xml'] = members['xl/worksheets/sheet1.xml'].replace(
        b'<c r="A1"', f'<c r="{coordinate}"'.encode())
    with ZipFile(path, 'w') as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    with pytest.raises(DomainError):
        inspect_file(path, 'sparse.xlsx')


def test_disguised_macro_part_is_rejected_by_content_type(tmp_path):
    path = tmp_path / 'report.docx'
    Document().save(path)
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members['[Content_Types].xml'] = members['[Content_Types].xml'].replace(b'</Types>',
        b'<Override PartName="/word/harmless.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
    members['word/harmless.bin'] = b'macro'
    with ZipFile(path, 'w') as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    with pytest.raises(DomainError):
        inspect_file(path, 'report.docx')


def test_locator_rows_are_stable_for_blank_spreadsheet_cells(tmp_path):
    path = tmp_path / 'blank.xlsx'
    workbook = Workbook()
    workbook.active.title = 'Inputs'
    workbook.active['C3'] = 42
    workbook.save(path)
    table = SpreadsheetParser().parse(path).tables[0]
    assert table.rows[2] == [None, None, 42]
    assert table.cell_locator(2, 2).model_dump(exclude_none=True) == {'sheet': 'Inputs', 'row': 3, 'column': 3}


def test_external_docx_hyperlink_never_triggers_network(tmp_path, monkeypatch):
    import socket
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    def disallow_network(*args, **kwargs):
        pytest.fail('Parser attempted outbound access')
    monkeypatch.setattr(socket.socket, 'connect', disallow_network)
    path = tmp_path / 'linked.docx'
    document = Document()
    paragraph = document.add_paragraph('Local text')
    rel_id = paragraph.part.relate_to('http://127.0.0.1:9/secret', RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rel_id)
    paragraph._p.append(link)
    document.save(path)
    assert DocxParser().parse(path).segments[0].text == 'Local text'
