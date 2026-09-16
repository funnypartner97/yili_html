from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from src.core.errors import DomainError
from src.files.parsers.base import inspect_file
from src.files.parsers.docx import DocxParser
from src.files.parsers.spreadsheet import SpreadsheetParser


def rewrite_member(path, name, transform):
    with ZipFile(path) as archive:
        members = {key: archive.read(key) for key in archive.namelist()}
    members[name] = transform(members[name])
    with ZipFile(path, 'w') as archive:
        for key, value in members.items():
            archive.writestr(key, value)


def spanned_table(document, *, width, rows, span, text='cell'):
    table = document.add_table(rows=0, cols=width)
    # Build physical XML cells, avoiding the expansion under test in fixtures.
    for _ in range(rows):
        row = OxmlElement('w:tr')
        cell = OxmlElement('w:tc')
        props = OxmlElement('w:tcPr')
        grid_span = OxmlElement('w:gridSpan')
        grid_span.set(qn('w:val'), str(span))
        props.append(grid_span)
        cell.append(props)
        paragraph = OxmlElement('w:p')
        run = OxmlElement('w:r')
        content = OxmlElement('w:t')
        content.text = text
        run.append(content)
        paragraph.append(run)
        cell.append(paragraph)
        row.append(cell)
        table._tbl.append(row)
    return table


@pytest.mark.parametrize('span', [100_000_000, 0, -1, 2])
def test_docx_preflight_rejects_excessive_or_inconsistent_grid_spans(tmp_path, span):
    path = tmp_path / 'span.docx'
    document = Document()
    spanned_table(document, width=1, rows=1, span=span)
    document.save(path)
    with pytest.raises(DomainError):
        inspect_file(path, path.name)


def test_docx_preflight_caps_cumulative_expanded_cells_across_tables(tmp_path):
    path = tmp_path / 'cumulative.docx'
    document = Document()
    for _ in range(2):
        spanned_table(document, width=256, rows=196, span=256)
    document.save(path)
    # Each table expands to 50,176 cells; together they exceed 100,000.
    with pytest.raises(DomainError) as error:
        inspect_file(path, path.name)
    assert error.value.code == 'source_limit_exceeded'


def test_docx_text_budget_is_checked_before_table_materialization(tmp_path, monkeypatch):
    import src.files.parsers.docx as parser
    path = tmp_path / 'text.docx'
    document = Document()
    spanned_table(document, width=2, rows=1, span=2, text='sixchars')
    document.save(path)
    monkeypatch.setattr(parser, 'MAX_TEXT', 10)
    def reject_construction(*args, **kwargs):
        pytest.fail('Table was constructed before enforcing the text budget')
    monkeypatch.setattr(parser, 'SourceTable', reject_construction)
    with pytest.raises(DomainError) as error:
        parser.DocxParser().parse(path)
    assert error.value.code == 'source_limit_exceeded'


def test_docx_valid_horizontal_and_vertical_merges_preserve_table_values(tmp_path):
    path = tmp_path / 'merged.docx'
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 1)).text = 'Revenue'
    document.save(path)
    assert DocxParser().parse(path).tables[0].rows == [['Revenue', 'Revenue'], ['Revenue', 'Revenue']]


@pytest.mark.parametrize('layout', ['table_only', 'leading', 'trailing', 'between'])
def test_docx_table_locator_resolves_to_actual_table_paragraph(tmp_path, layout):
    path = tmp_path / 'anchors.docx'
    document = Document()
    if layout in ('trailing', 'between'):
        document.add_paragraph('Before')
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Table evidence'
    if layout in ('leading', 'between'):
        document.add_paragraph('After')
    document.save(path)
    parsed = DocxParser().parse(path)
    locator = parsed.tables[0].locator
    resolved = [segment for segment in parsed.segments if segment.locator == locator]
    assert len(resolved) == 1
    assert resolved[0].text == 'Table evidence'
    # DOCX paragraph indices cover body and table-cell paragraphs in XML order.
    paragraphs = list(Document(path).element.body.iter(qn('w:p')))
    assert ''.join(paragraphs[locator.paragraph - 1].itertext()).count('Table evidence') >= 1
    assert [s.locator.paragraph for s in parsed.segments] == list(range(1, len(paragraphs) + 1))


@pytest.mark.parametrize('cached', [None, 6])
def test_xlsx_preserves_array_formula_text_range_and_cache(tmp_path, cached):
    path = tmp_path / 'array.xlsx'
    workbook = Workbook()
    workbook.active.title = 'Data'
    workbook.active['A1'] = 2
    workbook.active['A2'] = 3
    workbook.active['B1'] = ArrayFormula(ref='B1:B2', text='=A1:A2*3')
    workbook.save(path)
    if cached is not None:
        rewrite_member(path, 'xl/worksheets/sheet1.xml', lambda xml: xml.replace(
            b'</f><v></v>', b'</f><v>6</v>'))
    parsed = SpreadsheetParser().parse(path)
    formula = parsed.tables[0].formulas[0]
    assert formula.formula == '=A1:A2*3'
    assert formula.cached_value == cached
    assert formula.formula_kind == 'array'
    assert formula.formula_range == 'B1:B2'
    assert formula.locator.model_dump(exclude_none=True) == {'sheet': 'Data', 'row': 1, 'column': 2}
    assert parsed.tables[0].rows[0][1] == '=A1:A2*3'


def test_xlsx_preserves_data_table_formula_metadata_without_inventing_text(tmp_path):
    path = tmp_path / 'data-table.xlsx'
    workbook = Workbook()
    workbook.active['B1'] = DataTableFormula(ref='B1:B2', r1='A1')
    workbook.save(path)
    rewrite_member(path, 'xl/worksheets/sheet1.xml', lambda xml: xml.replace(
        b'</f><v></v>', b'</f><v>7</v>'))
    formula = SpreadsheetParser().parse(path).tables[0].formulas[0]
    assert formula.formula is None
    assert formula.formula_kind == 'dataTable'
    assert formula.formula_range == 'B1:B2'
    assert formula.attributes['r1'] == 'A1'
    assert formula.cached_value == 7
